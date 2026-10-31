"""The governed engine (api/_postgen) on its own: no database, fake model.

These carry over PostGenerator's red-line tests (tests/test_redlines.py,
test_eval_findings.py, test_template_filling.py there) to the merged engine,
plus what the merge itself changed: per-club settings instead of Chill Hi
High's hard-coded ones, an optional theme question, and the IG / Threads
variants rendered from the same content.
"""
import io
import re

import pytest
from pptx import Presentation

from api._postgen import context
from api._postgen.context import ClubProfile
from api._postgen.infrastructure.llm.fake import FakeLLMClient
from api._postgen.infrastructure.pptx.template_filler import fill_template
from api._postgen.modules.template_selector import pptx_registry
from api._postgen.renderers import social
from api._postgen.schemas.post_request import PostRequest
from api._postgen.service import GenerationService

MISSING = "[需由使用者提供]"

PROFILE = ClubProfile(
    club_id=1, name="測試分會",
    prefs={"register_url": "https://example.org/join", "brand_signature": "我們一起練表達。"},
    standing={"meeting_time": "19:10 ~ 21:00", "venue": "測試會議室", "transport": "捷運站步行 3 分鐘",
              "fee": "NTD150"},
)

PRE_EVENT = {
    "postType": "PRE_EVENT", "platforms": ["FACEBOOK", "PPTX_POSTER"],
    "topic": {"title_zh": "勇氣", "title_en": "Courage", "question_zh": "你是否曾鼓起勇氣跨出一步？"},
    "event": {"date": "2026-11-03", "venue": "社會創新實驗中心 107 會議室", "fee": "$100"},
    "host": {"name": "Spencer Ye", "role": "總主持人", "highlight": "自嘲式開場"},
}
POST_REVIEW = {
    "postType": "POST_REVIEW", "platforms": ["FACEBOOK", "PPTX_POSTER"],
    "topic": {"title_zh": "勇氣", "title_en": "Courage"},
    "meetingDate": "2026-11-03",
    "roles": [
        {"roleType": "總主持人", "name": "Spencer", "contribution": "自嘲式開場"},
        {"roleType": "國語演講員", "name": "庭敏", "contribution": "翻譯與游泳交織"},
        {"roleType": "講評員", "name": "Bill", "contribution": ""},
    ],
    "nextEvent": {"date": "2026-11-17"},
}


def run(raw, profile=PROFILE):
    svc = GenerationService(FakeLLMClient(), profile, key_source="fake", model="fake")
    resp = svc.generate(PostRequest.model_validate(raw))
    return svc, resp


def qc(resp, key):
    return next(c for c in resp.qc if c.key == key)


# ── red lines (PostGenerator tests/test_redlines.py) ────────────────

def test_missing_date_blocks():
    raw = {**PRE_EVENT, "event": {**PRE_EVENT["event"], "date": ""}}
    _, resp = run(raw)
    assert resp.blocked is True
    assert qc(resp, "facts").state == "crit" and "日期" in qc(resp, "facts").detail


def test_facts_are_verbatim_from_input():
    _, resp = run(PRE_EVENT)
    for value in ("社會創新實驗中心 107 會議室", "$100", "Spencer Ye", "https://example.org/join"):
        assert value in resp.facebook


def test_roles_must_match_input_verbatim():
    _, resp = run(POST_REVIEW)
    assert qc(resp, "roles").state == "ok"
    for name in ("Spencer", "庭敏", "Bill"):
        assert name in resp.facebook


def test_role_without_name_is_blocked():
    raw = {**POST_REVIEW, "roles": POST_REVIEW["roles"] + [{"roleType": "計時員", "name": ""}]}
    _, resp = run(raw)
    assert resp.blocked is True


def test_ai_written_highlights_are_disclosed():
    _, resp = run(POST_REVIEW)
    assert resp.ai_written_highlights == ["Bill"]
    assert any("Bill" in n.text and n.severity == "warn" for n in resp.compliance_notes)


def test_brand_signature_is_printed_verbatim():
    _, resp = run(POST_REVIEW)
    assert "我們一起練表達。" in resp.facebook


# ── what the merge changed ──────────────────────────────────────────

def test_weekday_is_computed_not_hard_coded():
    # 2026-11-05 is a Thursday; PostGenerator printed （二） for every date.
    raw = {**PRE_EVENT, "event": {**PRE_EVENT["event"], "date": "2026-11-05"}}
    _, resp = run(raw)
    assert "（四）" in resp.facebook and "（二）" not in resp.facebook


def test_club_standing_fills_blanks_and_is_disclosed():
    raw = {**PRE_EVENT, "event": {"date": "2026-11-03"}}
    _, resp = run(raw)
    assert "測試會議室" in resp.facebook and "19:10 ~ 21:00" in resp.facebook
    assert "捷運站步行 3 分鐘" in resp.facebook
    assert any("套用分會常態" in n.text for n in resp.compliance_notes)


def test_no_other_clubs_defaults_leak_in():
    _, resp = run(PRE_EVENT)
    for leaked in ("Chill Hi High", "forms.office.com", "19:15"):
        assert leaked not in resp.facebook


def test_the_theme_question_is_optional():
    raw = {**PRE_EVENT, "topic": {"title_zh": "勇氣", "title_en": "Courage"}}
    _, resp = run(raw)
    assert not resp.blocked
    assert "本週主題" not in resp.facebook and MISSING not in resp.facebook


NO_LINK = ClubProfile(club_id=2, name="沒設定的分會", standing=PROFILE.standing)


@pytest.mark.parametrize("raw", [PRE_EVENT, POST_REVIEW])
def test_the_register_link_is_optional(raw):
    svc, resp = run(raw, NO_LINK)
    assert not resp.blocked
    assert qc(resp, "cta").state == "ok" and qc(resp, "event").state == "ok"
    assert "報名連結" not in resp.facebook and "http" not in resp.facebook
    assert not resp.facebook.rstrip().endswith("：")
    assert not any("報名連結" in n.text for n in resp.compliance_notes)
    ctx, cc = svc.last_ctx, svc.last_cc
    assert "http" not in social.threads(ctx.req, cc, svc.settings)
    assert "個人簡介" not in social.instagram(ctx.req, cc, svc.settings, resp.facebook)
    if raw is POST_REVIEW:
        assert "👉 立即報名下一場例會\n" in resp.facebook


def test_an_unset_signature_is_flagged():
    _, resp = run(PRE_EVENT, NO_LINK)
    assert any("品牌定位句" in n.text for n in resp.compliance_notes)


def test_club_name_comes_from_the_profile():
    svc, resp = run(PRE_EVENT)
    assert "立即報名 測試分會" in resp.facebook
    assert "測試分會" in context.persona_head(PROFILE)


def test_prefs_are_validated():
    with pytest.raises(ValueError):
        context.clean_prefs({"post_persona": "字" * 1001})
    with pytest.raises(ValueError):
        context.clean_prefs({"post_default_length": "huge"})
    assert context.clean_prefs({"fixed_hashtags": "#A, #B，#C"}) == {"fixed_hashtags": ["#A", "#B", "#C"]}


def test_prompts_are_actually_loaded():
    # The merge once moved assets/ one level up and the red-line prompt silently
    # became empty. Reading it must fail loudly or contain the rules.
    svc = GenerationService(FakeLLMClient(), PROFILE)
    prompt = svc.story_prompt(svc.begin(PostRequest.model_validate(PRE_EVENT)))
    assert "三條不可違反的紅線" in prompt.system
    assert pptx_registry.list_templates([]), "built-in poster templates not found"


# ── platform variants ───────────────────────────────────────────────

def _variants(raw):
    svc, resp = run(raw)
    ctx, cc = svc.last_ctx, svc.last_cc
    return (resp, social.instagram(ctx.req, cc, svc.settings, resp.facebook),
            social.threads(ctx.req, cc, svc.settings))


def test_instagram_has_no_link_in_the_body():
    resp, ig, _ = _variants(PRE_EVENT)
    assert "https://example.org/join" in resp.facebook
    assert "https://example.org/join" not in ig and "個人簡介" in ig


@pytest.mark.parametrize("raw", [PRE_EVENT, POST_REVIEW])
def test_threads_fits_and_keeps_the_facts(raw):
    _, _, th = _variants(raw)
    assert len(th) <= social.THREADS_LIMIT
    assert "https://example.org/join" in th
    if raw is PRE_EVENT:
        assert "$100" in th and "11/3（二）" in th


# ── poster templates (PostGenerator tests/test_template_filling.py) ─

def _pptx_text(data: bytes) -> str:
    prs = Presentation(io.BytesIO(data))
    return "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)


@pytest.mark.parametrize("raw", [PRE_EVENT, POST_REVIEW])
def test_poster_has_no_leftover_tokens_and_the_facts(raw):
    _, resp = run(raw)
    plan = resp.poster_plan
    tpl = pptx_registry.resolve_from(pptx_registry.list_templates([], plan["use"]), plan["template_id"])
    text = _pptx_text(fill_template(pptx_registry.template_path(tpl), plan))
    assert not re.search(r"\{\{[A-Z0-9_]+\}\}", text)
    assert "我們一起練表達。" in text or raw is PRE_EVENT
    if raw is POST_REVIEW:
        assert "庭敏" in text


def test_photos_only_resolve_through_the_resolver():
    _, resp = run(PRE_EVENT)
    plan = {**resp.poster_plan, "photos": {"PHOTO_HOST": "https://elsewhere/x.jpg"}}
    tpl = pptx_registry.resolve_from(pptx_registry.list_templates([], "PRE_EVENT"), plan["template_id"])
    seen = []
    fill_template(pptx_registry.template_path(tpl), plan, resolve_photo=lambda ref: seen.append(ref))
    assert seen in ([], ["https://elsewhere/x.jpg"])     # asked, refused → placeholder kept
