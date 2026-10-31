"""Governed posts through the API, against a throwaway database.

Needs TEST_DATABASE_URL (see conftest.py); skipped without it. Each test runs
in its own club so they do not see each other's rows.
"""
import json
import os
import uuid

import pytest

pytestmark = pytest.mark.skipif(not os.environ.get("TEST_DATABASE_URL"),
                                reason="TEST_DATABASE_URL not set")


@pytest.fixture(scope="module")
def api():
    from fastapi.testclient import TestClient
    import api.index as index
    return index, TestClient(index.app)


@pytest.fixture()
def club(api):
    """A club with an officer, a member, this meeting and the next one."""
    index, _ = api
    tag = uuid.uuid4().hex[:8]
    with index.get_db() as conn:
        with conn.cursor() as cur:
            cur.execute("INSERT INTO clubs (name, name_zh, fee, settings) VALUES (%s,%s,%s,%s::jsonb)"
                        " RETURNING id", (f"Test {tag}", f"測試分會{tag}", "NTD150",
                                         json.dumps({"venue": "Venue: 測試會議室", "transit": "捷運 3 分鐘",
                                                     "timeRange": "19:10 ~ 21:00"})))
            cid = cur.fetchone()[0]
            users = {}
            for role in ("club_admin", "member"):
                name = f"{role}_{tag}"
                cur.execute("INSERT INTO users (username, password_hash, role, club_id, status)"
                            " VALUES (%s,'x',%s,%s,'active')", (name, role, cid))
                users[role] = name
            this = {"meetingDate": "2026-11-03", "meetingTheme": "勇氣 Courage", "meetingNo": "268",
                    "tme": "Spencer, DL5", "timer": "Eva", "ahCounter": "Frank",
                    "generalEvaluator": "Kelly", "speeches": [{"speaker": "庭敏", "title": "游泳"}],
                    "evaluators": ["Bill"], "venueInfo": ""}
            nxt = {"meetingDate": "2026-11-17", "timeRange": "19:00 ~ 21:00", "venueInfo": "下次的會議室"}
            ids = []
            for d in (this, nxt):
                cur.execute("INSERT INTO agendas (username, data, meeting_date, club_id)"
                            " VALUES (%s,%s::jsonb,%s,%s) RETURNING id",
                            (users["club_admin"], json.dumps(d), d["meetingDate"], cid))
                ids.append(cur.fetchone()[0])
            cur.execute("INSERT INTO club_post_prefs (club_id, prefs) VALUES (%s, %s::jsonb)",
                        (cid, json.dumps({"register_url": "https://example.org/join"})))
    headers = {r: {"Authorization": f"Bearer {index.make_token(u)}"} for r, u in users.items()}
    yield {"id": cid, "agenda": ids[0], "headers": headers}
    with index.get_db() as conn:
        with conn.cursor() as cur:
            cur.execute("DELETE FROM clubs WHERE id=%s", (cid,))
            cur.execute("DELETE FROM users WHERE username = ANY(%s)", (list(users.values()),))


def generate(client, club, kind="promo", **extra):
    body = {"club_id": club["id"], "agenda_id": club["agenda"], "kind": kind, **extra}
    r = client.post("/api/social-posts/generate", json=body, headers=club["headers"]["club_admin"])
    assert r.status_code == 200, r.text
    return r.json()["post"]


def test_promo_is_saved_pending_with_facts_from_the_database(api, club):
    _, client = api
    post = generate(client, club)
    assert post["reviewState"] == "pending"
    fb = post["variants"]["facebook"]["text"]
    for fact in ("11/3（二）", "19:10 ~ 21:00", "測試會議室", "NTD150", "Spencer"):
        assert fact in fb
    assert "Venue:" not in fb and "DL5" not in fb
    assert {"facebook", "instagram", "threads"} <= set(post["variants"])
    assert post["governance"]["qc"] and post["governance"]["blocked"] in (True, False)


def test_recap_lists_the_agenda_roles_and_the_next_meeting(api, club):
    _, client = api
    post = generate(client, club, kind="recap",
                    options={"role_notes": {"庭敏": "把游泳講成人生"}})
    fb = post["variants"]["facebook"]["text"]
    for name in ("Spencer", "庭敏", "Bill", "Kelly", "Eva", "Frank"):
        assert name in fb
    assert "11/17" in fb and "下次的會議室" in fb
    assert "庭敏" not in post["governance"]["aiWrittenHighlights"]
    assert "Bill" in post["governance"]["aiWrittenHighlights"]


def test_publish_and_poster_are_held_until_approved(api, club):
    index, client = api
    post = generate(client, club)
    with pytest.raises(index.HTTPException) as e:
        index._run_publish_job("x", club["id"], {"post_id": post["id"], "platforms": ["facebook"]})
    assert e.value.status_code == 403
    r = client.post(f"/api/social-posts/{post['id']}/poster", json={},
                    headers=club["headers"]["club_admin"])
    assert r.status_code == 403


def test_approval_needs_every_ai_written_role_confirmed(api, club):
    _, client = api
    post = generate(client, club, kind="recap")
    h = club["headers"]["club_admin"]
    names = post["governance"]["aiWrittenHighlights"]
    r = client.post(f"/api/social-posts/{post['id']}/review", json={"action": "approve"}, headers=h)
    assert r.status_code == 400 and names[0] in r.json()["detail"]
    r = client.post(f"/api/social-posts/{post['id']}/review",
                    json={"action": "approve", "confirmed": names}, headers=h)
    assert r.status_code == 200 and r.json()["reviewState"] == "approved"


def test_members_cannot_approve(api, club):
    _, client = api
    post = generate(client, club)
    r = client.post(f"/api/social-posts/{post['id']}/review", json={"action": "approve"},
                    headers=club["headers"]["member"])
    assert r.status_code == 403


def test_editing_approved_text_reopens_review(api, club):
    _, client = api
    h = club["headers"]["club_admin"]
    post = generate(client, club)
    client.post(f"/api/social-posts/{post['id']}/review", json={"action": "approve"}, headers=h)
    payload = {k: post[k] for k in ("title", "status", "body", "variants", "images", "kind")}
    payload["agenda_id"] = post["agendaId"]
    r = client.put(f"/api/social-posts/{post['id']}", json={**payload, "images": [{"url": "x"}]}, headers=h)
    assert r.json()["reviewState"] == "approved"            # images alone do not reopen
    payload["variants"]["threads"]["text"] += "（改一個字）"
    r = client.put(f"/api/social-posts/{post['id']}", json=payload, headers=h)
    assert r.json()["reviewState"] == "pending"


def test_blocked_drafts_cannot_be_approved(api, club):
    index, client = api
    with index.get_db() as conn:
        with conn.cursor() as cur:
            cur.execute("UPDATE agendas SET data = data || '{\"meetingTheme\": \"\"}'::jsonb WHERE id=%s",
                        (club["agenda"],))
    post = generate(client, club)
    assert post["governance"]["blocked"] is True
    r = client.post(f"/api/social-posts/{post['id']}/review", json={"action": "approve"},
                    headers=club["headers"]["club_admin"])
    assert r.status_code == 400


def test_regenerating_overwrites_the_same_post_and_reopens_review(api, club):
    _, client = api
    h = club["headers"]["club_admin"]
    post = generate(client, club)
    client.post(f"/api/social-posts/{post['id']}/review", json={"action": "approve"}, headers=h)
    again = generate(client, club, post_id=post["id"])
    assert again["id"] == post["id"] and again["reviewState"] == "pending"


def test_post_prefs_round_trip_and_feed_generation(api, club):
    _, client = api
    h = club["headers"]["club_admin"]
    r = client.put(f"/api/clubs/{club['id']}/post-prefs",
                   json={"brand_signature": "我們把話說清楚。", "register_url": "https://ex.org/j",
                         "fixed_hashtags": "#測試, #分會"}, headers=h)
    assert r.status_code == 200, r.text
    assert r.json()["prefs"]["fixed_hashtags"] == ["#測試", "#分會"]
    fb = generate(client, club, kind="recap")["variants"]["facebook"]["text"]
    assert "我們把話說清楚。" in fb and "https://ex.org/j" in fb and "#測試" in fb
    r = client.put(f"/api/clubs/{club['id']}/post-prefs", json={"post_persona": "字" * 1001}, headers=h)
    assert r.status_code == 400
    r = client.put(f"/api/clubs/{club['id']}/post-prefs", json={"brand_signature": "x"},
                   headers=club["headers"]["member"])
    assert r.status_code == 403


def test_other_kinds_still_use_the_free_writer(api, club):
    index, client = api
    body = {"club_id": club["id"], "agenda_id": club["agenda"], "kind": "other", "brief": "hi"}
    r = client.post("/api/social-posts/generate", json=body, headers=club["headers"]["club_admin"])
    # No key configured in tests: the free writer refuses rather than the engine running.
    assert r.status_code == 400 and "post" not in r.json()


def test_template_check_rejects_missing_tokens(tmp_path):
    from pptx import Presentation
    import api.index as index
    path = tmp_path / "t.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "{{HEADLINE}} {{DATE}} {{WHATEVER}}"
    prs.save(path)
    result = index.check_poster_template(str(path), "PRE_EVENT")
    assert any("VENUE" in e for e in result["errors"])
    assert any("WHATEVER" in w for w in result["warnings"])


class _FakeR2:
    """Stands in for R2: tests must never upload to the real bucket."""
    def __init__(self):
        self.objects = {}

    def put_object(self, Bucket, Key, Body, ContentType):
        self.objects[Key] = Body


def test_poster_after_approval_is_a_filled_pptx(api, club, monkeypatch):
    import io
    import re
    from pptx import Presentation
    index, client = api
    fake = _FakeR2()
    monkeypatch.setattr(index, "_r2", lambda: fake)
    h = club["headers"]["club_admin"]
    post = generate(client, club, kind="recap", options={"role_notes": {}})
    client.post(f"/api/social-posts/{post['id']}/review",
                json={"action": "approve", "confirmed": post["governance"]["aiWrittenHighlights"]},
                headers=h)
    r = client.post(f"/api/social-posts/{post['id']}/poster", json={}, headers=h)
    assert r.status_code == 200, r.text
    (key, data), = fake.objects.items()
    assert key.startswith(f"media/clubs/{club['id']}/social/") and key.endswith(".pptx")
    prs = Presentation(io.BytesIO(data))
    text = "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    assert "庭敏" in text and not re.search(r"\{\{[A-Z0-9_]+\}\}", text)
